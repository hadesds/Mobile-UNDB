import io
import shutil
import tempfile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from backend.leitura_curriculo import ArquivoInvalido, extrair, ler_arquivo
from backend.models import (
    ArquivoCurriculo,
    Competencia,
    CompetenciaDoPerfil,
    SinonimoCompetencia,
)

from .test_api import BaseAPITest

CURRICULO = """Ana Beatriz Souza
São Luís - MA | (98) 98877-6655 | ana.souza@email.com
linkedin.com/in/anabsouza | github.com/anabs

OBJETIVO
Estágio em desenvolvimento de software.

FORMAÇÃO ACADÊMICA
Bacharelado em Engenharia de Software - UNDB
Fev 2023 - Dez 2027 (previsão)
Cursando o 6º período

Técnico em Informática - IFMA
2019 - 2021

EXPERIÊNCIA PROFISSIONAL
Estagiária de Desenvolvimento Front-end | Maré Labs
03/2025 - Atual
• Desenvolvi telas em React Native com TypeScript.
• Integração com APIs REST em Django.

Monitora de Algoritmos - UNDB
Ago 2024 - Dez 2024
• Apoio a 40 alunos em Python.

PROJETOS
App de Caronas Universitárias - Projeto pessoal
2024
• App em Flutter com Firebase.

HABILIDADES
Python (avançado), JS, SQL, Git, Docker, Comunicação, Kotlin Multiplatform

IDIOMAS
Inglês - Avançado
Espanhol - Básico
Português - Nativo

CURSOS COMPLEMENTARES
React Native do Zero - Alura - 40h - 2024
Formação Python para Dados
Rocketseat - 2023 - 60 horas
"""


def _docx(paragrafos: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragrafos:
        doc.add_paragraph(p)
    tabela = doc.add_table(rows=1, cols=2)
    tabela.rows[0].cells[0].text = "Ferramentas"
    tabela.rows[0].cells[1].text = "Figma, Git"
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _pdf(linhas: list[str]) -> bytes:
    """PDF mínimo (Helvetica, WinAnsi) com uma linha de texto por linha."""

    def esc(t: str) -> str:
        return t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    conteudo = "BT /F1 11 Tf 14 TL 50 800 Td " + " ".join(f"({esc(l)}) Tj T*" for l in linhas) + " ET"
    stream = conteudo.encode("cp1252")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    saida = b"%PDF-1.4\n"
    posicoes = []
    for i, obj in enumerate(objetos, start=1):
        posicoes.append(len(saida))
        saida += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    saida += b"".join(b"%010d 00000 n \n" % p for p in posicoes)
    saida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref)
    return saida


class ExtracaoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        for nome, cat in [
            ("Python", "tecnica"), ("JavaScript", "tecnica"), ("SQL", "tecnica"), ("Git", "ferramenta"),
            ("Docker", "ferramenta"), ("React Native", "tecnica"), ("React", "tecnica"),
            ("TypeScript", "tecnica"), ("Django", "tecnica"), ("Flutter", "tecnica"),
            ("Comunicação", "comportamental"), ("Figma", "ferramenta"),
        ]:
            Competencia.objects.create(nome=nome, categoria=cat)
        SinonimoCompetencia.objects.create(competencia=Competencia.objects.get(nome="JavaScript"), termo="JS")

    def setUp(self):
        self.s = extrair(CURRICULO)

    def test_contato(self):
        c = self.s["contato"]
        self.assertEqual(c["email"], "ana.souza@email.com")
        self.assertEqual(c["telefone"], "(98) 98877-6655")
        self.assertIn("linkedin.com/in/anabsouza", c["linkedin"])
        self.assertIn("github.com/anabs", c["github"])

    def test_formacoes(self):
        f = self.s["formacoes"]
        self.assertEqual(len(f), 2, f)
        self.assertEqual(f[0]["nivel"], "graduacao")
        self.assertEqual(f[0]["curso"], "Engenharia de Software")
        self.assertEqual(f[0]["instituicao"], "UNDB")
        self.assertEqual(f[0]["inicio"], "2023-02-01")
        self.assertIsNone(f[0]["fim"])  # previsão no futuro = em andamento
        self.assertEqual(f[1]["nivel"], "tecnico")
        self.assertEqual(f[1]["curso"], "Informática")
        self.assertEqual(f[1]["instituicao"], "IFMA")
        self.assertEqual((f[1]["inicio"], f[1]["fim"]), ("2019-01-01", "2021-01-01"))
        self.assertEqual(f[1]["faltando"], [])

    def test_experiencias_e_projetos(self):
        e = self.s["experiencias"]
        self.assertEqual(len(e), 3, e)
        estagio, monitoria, projeto = e
        self.assertEqual(estagio["tipo"], "estagio")
        self.assertEqual(estagio["funcao"], "Estagiária de Desenvolvimento Front-end")
        self.assertEqual(estagio["organizacao"], "Maré Labs")
        self.assertEqual((estagio["inicio"], estagio["fim"]), ("2025-03-01", None))
        self.assertIn("React Native", estagio["descricao"])
        self.assertEqual(monitoria["tipo"], "academica")
        self.assertEqual(monitoria["organizacao"], "UNDB")
        self.assertEqual((monitoria["inicio"], monitoria["fim"]), ("2024-08-01", "2024-12-01"))
        self.assertEqual(projeto["tipo"], "projeto")
        self.assertEqual(projeto["funcao"], "App de Caronas Universitárias")

    def test_competencias_por_vocabulario_e_sinonimo(self):
        nomes = {c["competencia_nome"]: c for c in self.s["competencias"]}
        for nome in ["Python", "JavaScript", "SQL", "Git", "Docker", "Comunicação", "React Native",
                     "TypeScript", "Django", "Flutter"]:
            self.assertIn(nome, nomes)
        self.assertEqual(nomes["Python"]["nivel"], 3)
        self.assertEqual(nomes["Python"]["origem"], "secao")
        self.assertEqual(nomes["Flutter"]["origem"], "texto")
        # "React Native" não conta também como "React".
        self.assertNotIn("React", nomes)
        self.assertIn("Kotlin Multiplatform", self.s["termos_nao_reconhecidos"])

    def test_idiomas_sem_portugues(self):
        idiomas = {i["idioma"]: i for i in self.s["idiomas"]}
        self.assertEqual(set(idiomas), {"ingles", "espanhol"})
        self.assertEqual(idiomas["ingles"]["nivel"], 3)
        self.assertEqual(idiomas["espanhol"]["nivel"], 1)
        self.assertTrue(idiomas["ingles"]["nivel_informado"])

    def test_cursos_juntam_linhas_de_complemento(self):
        c = self.s["cursos"]
        self.assertEqual(len(c), 2, c)
        self.assertEqual(c[0]["nome"], "React Native do Zero")
        self.assertEqual(c[0]["instituicao"], "Alura")
        self.assertEqual(c[0]["carga_horaria"], 40)
        self.assertEqual(c[0]["concluido_em"], "2024-01-01")
        self.assertEqual(c[1]["nome"], "Formação Python para Dados")
        self.assertEqual(c[1]["instituicao"], "Rocketseat")
        self.assertEqual(c[1]["carga_horaria"], 60)

    def test_leitura_docx_inclui_tabelas_e_listas(self):
        leitura = ler_arquivo(_docx(["Habilidades", "Python, SQL"]), "cv.docx")
        self.assertEqual(leitura.tipo, "docx")
        self.assertIn("Figma, Git", leitura.texto)
        nomes = {c["competencia_nome"] for c in extrair(leitura.texto)["competencias"]}
        self.assertTrue({"Python", "SQL", "Figma", "Git"} <= nomes)

    def test_leitura_pdf(self):
        leitura = ler_arquivo(_pdf(["Ana Souza", "IDIOMAS", "Inglês - Fluente", "HABILIDADES", "Python, Docker"]), "cv.pdf")
        self.assertEqual(leitura.tipo, "pdf")
        self.assertEqual(leitura.paginas, 1)
        s = extrair(leitura.texto)
        self.assertEqual(s["idiomas"][0]["nivel"], 4)
        self.assertTrue({"Python", "Docker"} <= {c["competencia_nome"] for c in s["competencias"]})

    def test_arquivos_recusados(self):
        with self.assertRaisesMessage(ArquivoInvalido, ".doc"):
            ler_arquivo(b"\xd0\xcf\x11\xe0" + b"0" * 100, "cv.doc")
        with self.assertRaisesMessage(ArquivoInvalido, "PDF ou em Word"):
            ler_arquivo(b"ola mundo" * 20, "cv.txt")
        with self.assertRaisesMessage(ArquivoInvalido, "texto"):
            ler_arquivo(_pdf([""]), "vazio.pdf")


class ImportacaoAPITests(BaseAPITest):
    def setUp(self):
        self.midia = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.midia, ignore_errors=True)
        ajuste = override_settings(MEDIA_ROOT=self.midia)
        ajuste.enable()
        self.addCleanup(ajuste.disable)
        self.ana = self.novo_aluno()
        self.como(self.ana)

    def enviar(self, conteudo: bytes, nome="curriculo.docx"):
        return self.client.post(
            "/api/v1/me/perfil/arquivo/", {"arquivo": SimpleUploadedFile(nome, conteudo)}, format="multipart"
        )

    def test_envio_leitura_e_original(self):
        r = self.enviar(_docx(["Habilidades", "Python, SQL", "Idiomas", "Inglês intermediário"]))
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.data["arquivo"]["tipo"], "docx")
        self.assertEqual(r.data["arquivo"]["nome_original"], "curriculo.docx")
        self.assertIn("Python, SQL", r.data["texto"])
        self.assertEqual({c["competencia_nome"] for c in r.data["sugestoes"]["competencias"]}, {"Python", "SQL"})

        self.assertEqual(self.client.get("/api/v1/me/perfil/arquivo/").status_code, 200)
        self.assertEqual(self.client.get("/api/v1/me/perfil/").data["arquivo"]["nome_original"], "curriculo.docx")
        original = self.client.get("/api/v1/me/perfil/arquivo/original/")
        self.assertEqual(original.status_code, 200)
        self.assertTrue(b"".join(original.streaming_content).startswith(b"PK"))

    def test_novo_envio_substitui_e_apaga_o_anterior(self):
        self.enviar(_docx(["Habilidades", "Python"]))
        antigo = ArquivoCurriculo.objects.get().arquivo.path
        r = self.enviar(_pdf(["Ana Souza", "HABILIDADES", "SQL, Python e Docker"]), "cv.pdf")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(ArquivoCurriculo.objects.count(), 1)
        import os

        self.assertFalse(os.path.exists(antigo))
        self.assertEqual(self.client.delete("/api/v1/me/perfil/arquivo/").status_code, 204)
        self.assertEqual(self.client.get("/api/v1/me/perfil/arquivo/").status_code, 404)

    def test_arquivo_invalido(self):
        r = self.enviar(b"nada a ver" * 10, "foto.png")
        self.assertEqual(r.status_code, 400)
        self.assertIn("PDF", r.data["arquivo"][0])
        self.assertEqual(self.client.post("/api/v1/me/perfil/arquivo/", {}, format="multipart").status_code, 400)

    def test_sem_consentimento_e_outros_perfis(self):
        self.ana.aluno.registrar_consentimento(False)
        self.assertEqual(self.enviar(_docx(["Python"])).status_code, 403)
        self.como(self.gestor.usuario)
        self.assertEqual(self.client.get("/api/v1/me/perfil/arquivo/").status_code, 403)

    def test_importar_cria_itens_ignora_repetidos_e_aponta_erros(self):
        perfil = self.ana.aluno.perfil_competencias
        CompetenciaDoPerfil.objects.create(perfil=perfil, competencia=self.sql)
        corpo = {
            "formacoes": [{"nivel": "graduacao", "curso": "Engenharia de Software", "instituicao": "UNDB",
                           "inicio": "2023-02-01", "fim": None}],
            "experiencias": [
                {"tipo": "estagio", "funcao": "Estagiária", "organizacao": "Maré Labs", "descricao": "",
                 "inicio": "2025-03-01", "fim": None},
                {"tipo": "emprego", "funcao": "Sem lugar", "organizacao": "", "inicio": None, "fim": None},
            ],
            "competencias": [{"competencia": self.python.id, "nivel": 3}, {"competencia": self.sql.id, "nivel": None}],
            "idiomas": [{"idioma": "ingles", "nivel": 3}],
            "cursos": [{"nome": "React Native do Zero", "instituicao": "Alura", "carga_horaria": 40,
                        "concluido_em": "2024-01-01"}],
        }
        r = self.client.post("/api/v1/me/perfil/importar/", corpo, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.data["criados"], {"formacoes": 1, "experiencias": 1, "competencias": 1, "idiomas": 1, "cursos": 1})
        self.assertEqual(r.data["ignorados"], 1)
        self.assertEqual(len(r.data["erros"]), 1)
        self.assertEqual(r.data["erros"][0]["secao"], "experiencias")
        self.assertIn("organizacao", r.data["erros"][0]["erros"])
        self.assertEqual(r.data["perfil"]["completude"], 100)

        # Importar de novo não duplica nada.
        r = self.client.post("/api/v1/me/perfil/importar/", corpo, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(sum(r.data["criados"].values()), 0)

    def test_sugestoes_marcam_o_que_ja_existe(self):
        CompetenciaDoPerfil.objects.create(perfil=self.ana.aluno.perfil_competencias, competencia=self.python)
        r = self.enviar(_docx(["Habilidades", "Python, SQL"]))
        marcadas = {c["competencia_nome"]: c["ja_existe"] for c in r.data["sugestoes"]["competencias"]}
        self.assertEqual(marcadas, {"Python": True, "SQL": False})


class ExtracaoFormatosTests(TestCase):
    """Variações comuns: PDF sem linhas em branco, ano solitário, complemento de curso."""

    def test_sem_linhas_em_branco(self):
        texto = "\n".join(l for l in CURRICULO.splitlines() if l.strip())
        s = extrair(texto)
        tecnico = s["formacoes"][1]
        self.assertEqual(tecnico["curso"], "Informática")
        self.assertEqual((tecnico["inicio"], tecnico["fim"]), ("2019-01-01", "2021-01-01"))
        self.assertIsNone(s["formacoes"][0]["fim"])
        self.assertEqual(len(s["experiencias"]), 3)

    def test_ano_solitario_vira_o_ano_todo(self):
        projeto = extrair(CURRICULO)["experiencias"][2]
        self.assertEqual((projeto["inicio"], projeto["fim"]), ("2024-01-01", "2024-12-01"))
